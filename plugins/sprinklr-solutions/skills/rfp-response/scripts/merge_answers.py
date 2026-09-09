#!/usr/bin/env python3
"""Merge answers into the customer's files. The ONLY writer to source documents.

Subagents never do this — each owns one _work/answers/*.json and nothing else.
Single writer, no contention.

Nothing audit-related (confidence, sources, POC, tags, caveats) is ever written to a
customer file: those go to the customer. Audit data lives in STATE.md and the JSONs.

Usage:
    merge_answers.py                    dry run, prints every intended write
    merge_answers.py --apply            backup, then write
    merge_answers.py --apply --only b1 b2
"""
import argparse
import os
import shutil
import sys
from datetime import datetime

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rfp_lib import BASE, WORK, die, load_answers, load_config, sources  # noqa: E402


def backup(fname):
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    dest = os.path.join(WORK, "backups", f"{stamp} {os.path.basename(fname)}")
    os.makedirs(os.path.dirname(dest), exist_ok=True)
    shutil.copy2(os.path.join(BASE, fname), dest)
    return dest


def plan(doc, src):
    """Yield (row, col, value, label) for one batch."""
    w = src.get("write") or {}
    kind = src.get("answer_kind", "prose")
    verdicts = src.get("verdict_values", {"pass": "erfüllt", "fail": "nicht erfüllt"})
    for it in doc.get("items", []):
        row = it.get("row")
        if row is None:
            continue
        if kind == "criterion":
            if w.get("evidence_ref") and it.get("section"):
                tmpl = w.get("evidence_ref_template", "{section}")
                yield row, w["evidence_ref"], tmpl.format(section=it["section"]), "evidence ref"
            if w.get("verdict") and it.get("verdict"):
                yield row, w["verdict"], verdicts.get(it["verdict"], it["verdict"]), "verdict"
        else:
            if w.get("answer") and it.get("answer"):
                yield row, w["answer"], it["answer"], "answer"
            ful = it.get("std_fulfillment")
            cls = w.get("classification") or {}
            if ful and ful != "n/a" and cls:
                col = cls.get(ful)
                if col is None:
                    print(f"  !! {it['id']}: no column for {ful!r}", file=sys.stderr)
                else:
                    yield row, col, w.get("classification_mark", "x"), f"class:{ful}"
        for col, val in (w.get("constants") or {}).items():
            yield row, int(col), val, "constant"


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--apply", action="store_true", help="write (default: dry run)")
    ap.add_argument("--only", nargs="*", help="restrict to these batch ids")
    args = ap.parse_args()

    cfg = load_config()
    srcs = sources(cfg)
    _by_id, _by_batch, docs = load_answers()
    if not docs:
        die("no answer files found")

    per_file = {}
    for doc in docs:
        batch = doc["_batch"]
        if args.only and batch not in args.only:
            continue
        src = srcs.get(doc.get("source"))
        if not src:
            print(f"  !! {batch}: unknown source {doc.get('source')!r}", file=sys.stderr)
            continue
        if (src.get("write") or {}).get("mode") == "document":
            continue  # assembled into a document, not merged into cells
        per_file.setdefault(src["file"], (src, []))[1].extend(
            (batch, *wr) for wr in plan(doc, src))

    if not per_file:
        print("Nothing to merge.")
        return

    for fname, (src, writes) in per_file.items():
        never = {int(c) for c in (src.get("never_write") or [])}
        print(f"\n=== {fname}  ({len(writes)} cell writes)")
        seen = {}
        for batch, row, col, val, label in writes:
            if col in never:
                die(f"{fname}: batch {batch} tried to write column {col}, "
                    f"which is listed as never_write")
            key = (row, col)
            if key in seen and seen[key] != batch:
                die(f"{fname}: COLLISION r{row}c{col}: {seen[key]} vs {batch}")
            seen[key] = batch
            print(f"  [{batch}] r{row} c{col} {label:<16} = "
                  f"{' '.join(str(val).split())[:80]}")

        if not args.apply:
            continue

        import openpyxl
        dest = backup(fname)
        print(f"  backup -> {os.path.basename(dest)}")
        path = os.path.join(BASE, fname)
        wb = openpyxl.load_workbook(path)  # keep formulas
        ws = wb[src.get("sheet") or wb.sheetnames[0]]
        for _b, row, col, val, _l in writes:
            ws.cell(row, col).value = val
        wb.save(path)
        print(f"  WRITTEN: {fname}")

    if not args.apply:
        print("\nDry run only. Re-run with --apply to write.")


if __name__ == "__main__":
    main()
