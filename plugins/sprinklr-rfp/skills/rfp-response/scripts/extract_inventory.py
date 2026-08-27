#!/usr/bin/env python3
"""Extract in-scope questions from the customer's files into _work/inventory/*.json.

Read-only against source files. Agents consume the inventory and never open a
customer document themselves.

Handles xlsx and csv. For other formats, convert first or have a subagent write the
inventory JSON directly in the same shape (see references/config-schema.md).
"""
import csv
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from rfp_lib import BASE, WORK, load_config, die  # noqa: E402


def cell(ws, row, col):
    v = ws.cell(row, col).value
    return v if v != "" else None


def extract_prose(ws, src):
    """Rows where section_col is set are headers; rows with a question are items."""
    lo, hi = src["row_range"]
    sec_col = src.get("section_col")
    q_col = src["question_col"]
    prefix = src.get("id_prefix", src["id"].upper())
    out, section, section_title = [], None, None
    for r in range(lo, hi + 1):
        sec = cell(ws, r, sec_col) if sec_col else None
        q = cell(ws, r, q_col)
        if sec:
            section, section_title = str(sec).strip(), str(q or "").strip()
            continue
        if not q:
            continue
        extra = {k: cell(ws, r, c) for k, c in (src.get("read_cols") or {}).items()}
        out.append({
            "id": f"{prefix}-{section}-r{r}" if section else f"{prefix}-r{r}",
            "source": src["id"], "file": src["file"], "row": r,
            "section": section, "section_title": section_title,
            "question": str(q).strip(), "extra": extra,
        })
    return out


def extract_criterion(ws, src):
    id_col = src["id_col"]
    first = src.get("first_data_row", 2)
    last = src.get("last_data_row") or ws.max_row
    out = []
    for r in range(first, last + 1):
        gid = cell(ws, r, id_col)
        if not gid:
            continue
        extra = {k: cell(ws, r, c) for k, c in (src.get("read_cols") or {}).items()}
        out.append({
            "id": str(gid).strip(), "source": src["id"], "file": src["file"],
            "row": r, "section": extra.get("category"),
            "section_title": extra.get("title"),
            "question": extra.get("description") or extra.get("title") or "",
            "extra": extra,
        })
    return out


def extract_csv(src):
    path = os.path.join(BASE, src["file"])
    prefix = src.get("id_prefix", src["id"].upper())
    out = []
    with open(path, newline="", encoding=src.get("encoding", "utf-8-sig")) as fh:
        for n, row in enumerate(csv.DictReader(fh), start=2):
            q = row.get(src["question_col"])
            if not q or not q.strip():
                continue
            idv = row.get(src.get("id_col") or "") or f"{prefix}-r{n}"
            out.append({
                "id": str(idv).strip(), "source": src["id"], "file": src["file"],
                "row": n, "section": row.get(src.get("section_col") or ""),
                "section_title": None, "question": q.strip(),
                "extra": {k: row.get(c) for k, c in (src.get("read_cols") or {}).items()},
            })
    return out


def main():
    cfg = load_config()
    os.makedirs(os.path.join(WORK, "inventory"), exist_ok=True)
    total = 0
    for src in cfg.get("sources", []):
        if (src.get("write") or {}).get("mode") == "document" and src.get("skip_extract"):
            continue
        kind = src.get("type", "xlsx")
        if kind == "csv":
            rows = extract_csv(src)
        elif kind == "xlsx":
            try:
                import openpyxl
            except ImportError:
                die("openpyxl not installed — pip install openpyxl "
                    "(use a venv if the system Python is externally managed)")
            path = os.path.join(BASE, src["file"])
            if not os.path.exists(path):
                die(f"source file missing: {path}")
            wb = openpyxl.load_workbook(path, data_only=True)
            sheet = src.get("sheet") or wb.sheetnames[0]
            if sheet not in wb.sheetnames:
                die(f"{src['file']}: no sheet {sheet!r} (have {wb.sheetnames})")
            ws = wb[sheet]
            rows = (extract_criterion(ws, src) if src.get("answer_kind") == "criterion"
                    else extract_prose(ws, src))
        else:
            die(f"{src['id']}: unsupported type {kind!r} — convert it, or write "
                f"_work/inventory/{src['id']}.json by hand")

        seen = set()
        for r in rows:
            if r["id"] in seen:
                die(f"{src['id']}: duplicate item id {r['id']!r} — ids must be unique")
            seen.add(r["id"])

        out = os.path.join(WORK, "inventory", f"{src['id']}.json")
        with open(out, "w") as fh:
            json.dump(rows, fh, ensure_ascii=False, indent=2)
        print(f"{src['id']}: {len(rows)} items -> _work/inventory/{src['id']}.json")
        total += len(rows)
    print(f"\nTotal in scope: {total}")


if __name__ == "__main__":
    main()
